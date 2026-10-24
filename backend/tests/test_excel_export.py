"""The Excel export is a live model: formulas linked to named Assumptions, no hard-coded results in calculation tabs."""
import io
from copy import deepcopy

from openpyxl import load_workbook

from app.excel_export import export_project
from engine.model import run_model
from seed.demo_cases import CASES


def _wb(case):
    proj = {"id": 1, "name": case["profile"]["name"], "profile": case["profile"], "assumptions": deepcopy(case["assumptions"]), "meta": {}}
    res = run_model(proj["profile"], proj["assumptions"])
    return load_workbook(io.BytesIO(export_project(proj, res, None))), res


def test_workbook_structure_and_named_inputs():
    wb, _ = _wb(CASES[0])
    for tab in ("Cover", "Outputs", "Assumptions", "Development", "Construction", "Sales", "Cash Flow", "Waterfall", "Budget", "Sensitivity"):
        assert tab in wb.sheetnames
    names = set(wb.defined_names.keys())
    assert {"in_site_land_price", "in_cost_hard_per_area", "in_rev_absorption_per_month", "TotalCost", "M_Lev", "M_Unlev"} <= names
    assert wb.calculation.iterate  # debt sizing is circular by design
    assert len(wb["Outputs"]._charts) >= 4


def test_calculation_tabs_contain_formulas_not_values():
    wb, _ = _wb(CASES[0])
    for tab in ("Construction", "Sales", "Cash Flow", "Waterfall"):
        ws = wb[tab]
        numeric_hardcodes = []
        for row in ws.iter_rows(min_row=5, min_col=4, max_col=40):
            for c in row:
                if isinstance(c.value, (int, float)) and not (c.row == 4):
                    numeric_hardcodes.append(c.coordinate)
        assert not numeric_hardcodes, f"{tab} has hard-coded numbers at {numeric_hardcodes[:5]}"
    out = wb["Outputs"]
    assert str(out["B6"].value).startswith("=")  # KPIs are formulas


def test_assumptions_carry_the_project_values():
    wb, _ = _wb(CASES[0])
    ref = wb.defined_names["in_site_land_price"].attr_text.split("!")[1].replace("$", "")
    assert wb["Assumptions"][ref].value == CASES[0]["assumptions"]["site.land_price"]
    ref = wb.defined_names["in_fin_senior_ltc_pct"].attr_text.split("!")[1].replace("$", "")
    assert abs(wb["Assumptions"][ref].value - CASES[0]["assumptions"]["fin.senior_ltc_pct"] / 100) < 1e-12  # percent as fraction


def test_rental_case_exports_rental_tab():
    wb, _ = _wb(CASES[2])
    assert wb["Rental"].sheet_state == "visible"


def test_monthly_charts_start_at_month_zero_and_columns_fit_values():
    """Regression: monthly chart series must start at column D (month 0), not pull in the Total/note columns,
    the single-series units chart must not vary colors per month, and month columns must fit 9-digit negatives."""
    import re
    import zipfile
    case = CASES[0]
    proj = {"id": 1, "name": case["profile"]["name"], "profile": case["profile"], "assumptions": deepcopy(case["assumptions"]), "meta": {}}
    data = export_project(proj, run_model(proj["profile"], proj["assumptions"]), None)
    z = zipfile.ZipFile(io.BytesIO(data))
    charts = {n: z.read(n).decode() for n in z.namelist() if n.startswith("xl/charts/chart")}
    monthly = {n: x for n, x in charts.items() if "'Cash Flow'!$" in x or "'Sales'!$" in x or "Sales!$" in x}
    assert monthly, "expected monthly charts"
    checked = 0
    for name, xml in monthly.items():
        for ref in re.findall(r"<(?:c:)?val>\s*<(?:c:)?numRef>\s*<(?:c:)?f>([^<]+)</", xml):
            assert re.search(r"!\$D\$\d+:", ref), f"{name}: series values {ref} should start at column D"
            checked += 1
    assert checked >= 4  # 3 cumulative lines + units sold
    units = next(x for x in charts.values() if "Units sold per month" in x)
    assert re.search(r'<(?:c:)?varyColors val="0"', units) and not re.search(r"<(?:c:)?legend>", units)
    wb = load_workbook(io.BytesIO(data))
    assert wb["Cash Flow"].column_dimensions["D"].width >= 14
    assert wb["Development"].column_dimensions["D"].width >= 16 and wb["Sensitivity"].column_dimensions["C"].width >= 16
